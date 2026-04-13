from io import BytesIO

from openpyxl import Workbook

from app.services.workbook_ingestion import WorkbookIngestionService
from app.utils.file_validation import ValidatedUpload


def build_workbook_bytes(sheet_populators: dict[str, list[tuple[str, object]]]) -> bytes:
    workbook = Workbook()
    first_sheet = True
    for sheet_name, cells in sheet_populators.items():
        if first_sheet:
            sheet = workbook.active
            sheet.title = sheet_name
            first_sheet = False
        else:
            sheet = workbook.create_sheet(title=sheet_name)

        for cell_reference, value in cells:
            sheet[cell_reference] = value

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_manifest_contains_normalization_and_profile_metadata() -> None:
    workbook_bytes = build_workbook_bytes(
        {
            "Sales": [
                ("A1", "Product"),
                ("B1", "2024"),
                ("C1", "2025"),
                ("A2", "Alpha"),
                ("B2", 10),
                ("C2", 12),
                ("A3", "Beta"),
                ("B3", 20),
                ("C3", 22),
            ],
            "Events": [
                ("A1", "Date"),
                ("B1", "Owner"),
                ("C1", "Notes"),
                ("A2", "2026-01-01"),
                ("B2", "Platform"),
                ("C2", "Kickoff review"),
                ("A3", "2026-01-02"),
                ("B3", "Architecture"),
                ("C3", "Migration workshop"),
            ],
        }
    )
    service = WorkbookIngestionService()

    manifest = service.build_manifest(
        upload_id="upl_profile",
        validated_upload=ValidatedUpload(
            file_name="report.xlsx",
            safe_file_name="report.xlsx",
            file_type="xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=len(workbook_bytes),
        ),
        file_bytes=workbook_bytes,
    )

    sales_table = next(table for table in manifest.tables if table.sheet_id == "sheet_01")
    events_table = next(table for table in manifest.tables if table.sheet_id == "sheet_02")

    assert sales_table.orientation == "wide_form"
    assert sales_table.normalization.status == "melted"
    assert sales_table.normalization.output_columns == ["Product", "variable", "value"]
    assert sales_table.source_reference is not None
    assert sales_table.stats.row_count == 4
    assert any(column.role == "numeric" for column in sales_table.columns)

    assert events_table.orientation == "long_form"
    assert events_table.normalization.status == "none"
    assert any(column.role == "datetime" for column in events_table.columns)
    assert events_table.stats.primary_mode in {"table", "summary", "chart"}
