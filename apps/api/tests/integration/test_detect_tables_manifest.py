from io import BytesIO

from openpyxl import Workbook

from app.services.workbook_ingestion import WorkbookIngestionService
from app.utils.file_validation import ValidatedUpload


def build_workbook_bytes(sheet_populators: dict[str, list[tuple[str, object]]]) -> bytes:
    workbook = Workbook()
    first_sheet = True
    for sheet_name, cells in sheet_populators.items():
        if first_sheet:
            sheet = workbook.active
            sheet.title = sheet_name
            first_sheet = False
        else:
            sheet = workbook.create_sheet(title=sheet_name)

        for cell_reference, value in cells:
            sheet[cell_reference] = value

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_manifest_contains_multiple_tables_for_multi_block_sheet() -> None:
    workbook_bytes = build_workbook_bytes(
        {
            "Report": [
                ("A1", "Product"),
                ("B1", "Revenue"),
                ("A2", "Alpha"),
                ("B2", 10),
                ("A3", "Beta"),
                ("B3", 20),
                ("A5", "Team"),
                ("B5", "Count"),
                ("A6", "Platform"),
                ("B6", 4),
            ]
        }
    )
    service = WorkbookIngestionService()

    manifest = service.build_manifest(
        upload_id="upl_blocks",
        validated_upload=ValidatedUpload(
            file_name="report.xlsx",
            safe_file_name="report.xlsx",
            file_type="xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=len(workbook_bytes),
        ),
        file_bytes=workbook_bytes,
    )

    assert manifest.workbook.table_count == 2
    assert len(manifest.tables) == 2
    assert manifest.tables[0].bounds.start_row == 1
    assert manifest.tables[0].bounds.end_row == 3
    assert manifest.tables[1].bounds.start_row == 5
    assert manifest.tables[1].bounds.end_row == 6
    assert all(table.review_required is False for table in manifest.tables)


def test_manifest_marks_repeated_header_layouts_for_review() -> None:
    workbook_bytes = build_workbook_bytes(
        {
            "Paged": [
                ("A1", "Product"),
                ("B1", "Revenue"),
                ("A2", "Alpha"),
                ("B2", 10),
                ("A3", "Beta"),
                ("B3", 20),
                ("A4", "Product"),
                ("B4", "Revenue"),
                ("A5", "Gamma"),
                ("B5", 30),
            ]
        }
    )
    service = WorkbookIngestionService()

    manifest = service.build_manifest(
        upload_id="upl_review",
        validated_upload=ValidatedUpload(
            file_name="repeated.xlsx",
            safe_file_name="repeated.xlsx",
            file_type="xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=len(workbook_bytes),
        ),
        file_bytes=workbook_bytes,
    )

    assert manifest.workbook.table_count == 1
    assert manifest.tables[0].review_required is True
    assert any(
        "repeated header" in reason.lower()
        for reason in manifest.tables[0].detection_reasons
    )
